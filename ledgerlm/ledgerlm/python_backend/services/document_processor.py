#!/usr/bin/env python3
"""
Multi-Format Document Processor for LedgerLM
Handles PDF, Excel, CSV, text, image, and web document formats
"""

import os
import tempfile
import logging
from typing import List, Dict, Any, Optional, Union
import json
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
import base64
import io

# Import processing libraries
import pandas as pd
from PIL import Image
import pytesseract

# Import existing processors
from parsers.pdf_processor import process_pdf
from services.company_extractor import extract_company_info

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("multi-format-processor")

class HTMLTextExtractor(HTMLParser):
    """Extract text content from HTML files"""
    
    def __init__(self):
        super().__init__()
        self.text_content = []
        self.in_script = False
        self.in_style = False
    
    def handle_starttag(self, tag, attrs):
        if tag.lower() in ['script', 'style']:
            self.in_script = True
            self.in_style = True
    
    def handle_endtag(self, tag):
        if tag.lower() in ['script', 'style']:
            self.in_script = False
            self.in_style = False
    
    def handle_data(self, data):
        if not (self.in_script or self.in_style):
            clean_data = data.strip()
            if clean_data:
                self.text_content.append(clean_data)
    
    def get_text(self):
        return '\n'.join(self.text_content)

def _create_enhanced_table_text(df, sheet_name, include_time_series_analysis=True):
    """Create enhanced table text with proper formatting and comprehensive financial analysis"""
    try:
        if df.empty:
            return f"Sheet: {sheet_name}\nNo data available"
        
        # Create formatted table text
        table_text = f"Sheet: {sheet_name}\n"
        table_text += "=" * 50 + "\n"
        
        # Add column headers
        headers = " | ".join(str(col).ljust(15) for col in df.columns)
        table_text += headers + "\n"
        table_text += "-" * len(headers) + "\n"
        
        # Add data rows (limit to first 100 rows for performance)
        display_rows = min(100, len(df))
        for idx in range(display_rows):
            row = df.iloc[idx]
            row_text = " | ".join(str(val).ljust(15)[:15] for val in row.values)
            table_text += row_text + "\n"
        
        if len(df) > display_rows:
            table_text += f"\n... and {len(df) - display_rows} more rows\n"
        
        # Enhanced financial analysis with time-series intelligence
        if include_time_series_analysis:
            try:
                # Import time-series analyzer (lazy import)
                from time_series_analyzer import analyze_financial_time_series
                
                # Perform time-series analysis
                ts_analysis = analyze_financial_time_series(
                    df, 
                    metadata={'sheet_name': sheet_name, 'source': 'excel_processing'}
                )
                
                # Add time-series analysis results to table text
                table_text += _format_time_series_analysis_text(ts_analysis)
                
            except ImportError:
                logger.warning("Time-series analyzer not available, using basic analysis")
                table_text += _create_basic_financial_summary(df)
            except Exception as e:
                logger.warning(f"Time-series analysis failed for {sheet_name}: {e}")
                table_text += _create_basic_financial_summary(df)
        else:
            table_text += _create_basic_financial_summary(df)
        
        return table_text
        
    except Exception as e:
        logger.warning(f"Error creating enhanced table text: {e}")
        return f"Sheet: {sheet_name}\n{df.to_string(index=False)}"

def _format_time_series_analysis_text(ts_analysis):
    """Format time-series analysis results as text"""
    try:
        analysis_text = "\n" + "=" * 50 + "\n"
        analysis_text += "ADVANCED TIME-SERIES ANALYSIS:\n"
        analysis_text += "-" * 35 + "\n"
        
        # Fiscal calendar information
        fiscal_info = ts_analysis.get('fiscal_calendar', {})
        if fiscal_info:
            analysis_text += f"Fiscal Calendar: {fiscal_info.get('type', 'Unknown')}\n"
            if fiscal_info.get('start_month', 1) != 1:
                analysis_text += f"Fiscal Year Starts: Month {fiscal_info['start_month']}\n"
            analysis_text += f"Detection Confidence: {fiscal_info.get('confidence', 0):.1%}\n\n"
        
        # Data normalization summary
        normalization = ts_analysis.get('data_normalization', {})
        if normalization:
            analysis_text += "Data Structure Analysis:\n"
            analysis_text += f"  Total Records: {len(normalization.get('canonical_data', []))}\n"
            analysis_text += f"  Data Quality Score: {normalization.get('data_quality_score', 0):.1%}\n"
            
            # Detected dimensions
            entities = normalization.get('entities', [])
            if entities:
                analysis_text += f"  Entities: {', '.join(entities[:3])}{'...' if len(entities) > 3 else ''}\n"
            
            metrics = normalization.get('metrics', [])
            if metrics:
                analysis_text += f"  Metrics: {', '.join(metrics[:3])}{'...' if len(metrics) > 3 else ''}\n"
            
            scenarios = normalization.get('scenarios', [])
            if scenarios:
                analysis_text += f"  Scenarios: {', '.join(scenarios[:3])}{'...' if len(scenarios) > 3 else ''}\n"
            
            analysis_text += "\n"
        
        # Financial metrics insights
        financial_metrics = ts_analysis.get('financial_metrics', {})
        if financial_metrics:
            # Key insights
            insights = financial_metrics.get('key_insights', [])
            if insights:
                analysis_text += "Key Financial Insights:\n"
                for insight in insights[:5]:  # Top 5 insights
                    analysis_text += f"  • {insight}\n"
                analysis_text += "\n"
            
            # Growth rates summary
            growth_rates = financial_metrics.get('growth_rates', {})
            if growth_rates:
                analysis_text += "Growth Rate Analysis:\n"
                for metric, growth_data in list(growth_rates.items())[:3]:  # Top 3 metrics
                    if growth_data.get('cagr'):
                        analysis_text += f"  {metric}: CAGR {growth_data['cagr']}%"
                        if growth_data.get('average_growth'):
                            analysis_text += f", Avg Growth: {growth_data['average_growth']}%"
                        analysis_text += "\n"
                analysis_text += "\n"
            
            # Variance analysis
            variance_analysis = financial_metrics.get('variance_analysis', {})
            if variance_analysis:
                analysis_text += "Variance Analysis:\n"
                for metric, variance_data in list(variance_analysis.items())[:3]:
                    if variance_data.get('vs_budget'):
                        var_info = variance_data['vs_budget']
                        analysis_text += f"  {metric} vs Budget: {var_info['percentage_variance']:+.1f}% ({var_info['performance']})\n"
                    elif variance_data.get('vs_forecast'):
                        var_info = variance_data['vs_forecast']
                        analysis_text += f"  {metric} vs Forecast: {var_info['percentage_variance']:+.1f}% ({var_info['performance']})\n"
                analysis_text += "\n"
        
        # Overall summary
        summary = ts_analysis.get('summary', {})
        if summary:
            analysis_text += f"Analysis Summary:\n"
            analysis_text += f"  Analysis Date: {summary.get('analysis_date', 'N/A')}\n"
            analysis_text += f"  Total Records Processed: {summary.get('total_records', 0)}\n"
            analysis_text += f"  Overall Data Quality: {summary.get('data_quality_score', 0):.1%}\n\n"
        
        return analysis_text
        
    except Exception as e:
        logger.warning(f"Error formatting time-series analysis: {e}")
        return "\n" + "=" * 50 + "\nTime-series analysis formatting error\n\n"

def _create_basic_financial_summary(df):
    """Create basic financial summary when time-series analysis is not available"""
    try:
        # Add summary statistics for numeric columns
        numeric_cols = df.select_dtypes(include=['number']).columns
        if len(numeric_cols) > 0:
            summary_text = "\n" + "=" * 50 + "\n"
            summary_text += "BASIC FINANCIAL SUMMARY:\n"
            summary_text += "-" * 25 + "\n"
            
            for col in numeric_cols:
                if df[col].notna().any():
                    summary_text += f"{col}:\n"
                    summary_text += f"  Total: {df[col].sum():,.2f}\n"
                    summary_text += f"  Average: {df[col].mean():,.2f}\n"
                    summary_text += f"  Min: {df[col].min():,.2f}\n"
                    summary_text += f"  Max: {df[col].max():,.2f}\n\n"
            
            return summary_text
        
        return "\n"
        
    except Exception as e:
        logger.warning(f"Error creating basic financial summary: {e}")
        return "\n"

def get_file_format(filename: str) -> str:
    """Get standardized file format from filename"""
    ext = filename.lower().split('.')[-1] if '.' in filename else ''
    
    format_mapping = {
        'pdf': 'pdf',
        'xlsx': 'xlsx', 
        'xls': 'xls',
        'csv': 'csv',
        'txt': 'txt',
        'md': 'md',
        'json': 'json',
        'xml': 'xml',
        'html': 'html',
        'htm': 'html',
        'jpg': 'image',
        'jpeg': 'image',
        'png': 'image',
        'bmp': 'image',
        'gif': 'image',
        'docx': 'docx',
        'doc': 'docx',
    }
    
    return format_mapping.get(ext, 'unknown')

def validate_file_format(filename: str) -> bool:
    """Check if file format is supported"""
    file_format = get_file_format(filename)
    supported_formats = get_supported_formats()
    
    # For images, we need to check the actual extension
    if file_format == 'image':
        ext = filename.lower().split('.')[-1] if '.' in filename else ''
        return ext in ['jpg', 'jpeg', 'png', 'bmp', 'gif']
    
    return file_format in [key for key in supported_formats.keys() if key != 'image']

def process_excel_file(file_path: str) -> List[Dict[str, Any]]:
    """Process Excel files (xlsx/xls) and extract structured data"""
    
    # Check file size before processing (increased limit for larger Excel files)
    # Supports up to 2GB for very large financial files with millions of rows
    MAX_EXCEL_SIZE = 2000 * 1024 * 1024  # 2GB for very large financial files
    try:
        file_size = os.path.getsize(file_path)
        if file_size > MAX_EXCEL_SIZE:
            size_mb = file_size / (1024 * 1024)
            logger.error(f"Excel file too large: {size_mb:.1f}MB (max: {MAX_EXCEL_SIZE/(1024*1024)}MB)")
            return [{'text': f"Error: Excel file too large ({size_mb:.1f}MB). Maximum supported size is {MAX_EXCEL_SIZE/(1024*1024)}MB.", 'error': True}]
        
        logger.info(f"Processing Excel file: {file_size/(1024*1024):.1f}MB")
    except Exception as e:
        logger.warning(f"Could not check file size: {e}")
    
    def try_with_engine(engine=None):
        """Helper function to try reading Excel with specific engine"""
        if engine:
            excel_file = pd.ExcelFile(file_path, engine=engine)
        else:
            excel_file = pd.ExcelFile(file_path)  # Use default behavior
        
        content = []
        for sheet_name in excel_file.sheet_names:
            if engine:
                df = pd.read_excel(file_path, sheet_name=sheet_name, engine=engine)
            else:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
            
            # Enhanced data extraction for comprehensive analysis
            sheet_text = f"Sheet: {sheet_name}\n"
            sheet_text += df.to_string(index=False)
            
            # Create comprehensive financial data structure with time-series analysis
            enhanced_table_text = _create_enhanced_table_text(df, sheet_name)
            
            # Perform time-series analysis for metadata
            time_series_analysis = None
            try:
                from time_series_analyzer import analyze_financial_time_series
                time_series_analysis = analyze_financial_time_series(
                    df, 
                    metadata={'sheet_name': sheet_name, 'source': 'excel_file', 'file_path': file_path}
                )
                logger.info(f"Time-series analysis completed for sheet: {sheet_name}")
            except ImportError:
                logger.info("Time-series analyzer not available for Excel processing")
            except Exception as e:
                logger.warning(f"Time-series analysis failed for {sheet_name}: {e}")
            
            sheet_data = {
                'sheet_name': sheet_name,
                'text': sheet_text,
                'structured_data': df.to_dict('records') if len(df) > 0 else [],
                'columns': list(df.columns),
                'row_count': len(df),
                'total_text': enhanced_table_text,
                'time_series_analysis': time_series_analysis,  # New: Include time-series analysis
                'data_quality_score': time_series_analysis.get('summary', {}).get('data_quality_score', 0.8) if time_series_analysis else 0.8,
                'detected_scenarios': time_series_analysis.get('data_normalization', {}).get('scenarios', []) if time_series_analysis else [],
                'detected_metrics': time_series_analysis.get('data_normalization', {}).get('metrics', []) if time_series_analysis else [],
                'fiscal_calendar': time_series_analysis.get('fiscal_calendar', {}) if time_series_analysis else {},
                'has_time_series_data': time_series_analysis is not None,
                'pages': 0,  # Excel sheets don't have pages
                'chunks': 0,  # Will be set during chunking
                'file_size': os.path.getsize(file_path) if os.path.exists(file_path) else 0
            }
            
            content.append(sheet_data)
        
        return content
    
    # List of engines to try in order of preference
    engines_to_try = [None, 'openpyxl', 'xlrd', 'calamine']
    
    def _try_csv_fallback(file_path: str) -> List[Dict[str, Any]] | None:
        """Try to process file as CSV if it appears to be text-based content"""
        try:
            with open(file_path, 'rb') as f:
                file_header = f.read(1000)  # Read first 1000 bytes for better detection
            
            # Check if it's an HTML error page
            if file_header.startswith(b'<!DOCTYPE') or file_header.startswith(b'<html') or file_header.startswith(b'<HTML'):
                logger.error(f"Downloaded file appears to be HTML, not Excel: {file_path}")
                return [{'text': "Error: Downloaded file is an HTML page, not an Excel file. Check source permissions.", 'error': True}]
            
            # Check for ZIP signature (Excel files are ZIP archives)
            # PK\x03\x04 is the ZIP magic number
            if file_header.startswith(b'PK\x03\x04'):
                return None  # It's a valid ZIP/Excel file, don't try CSV fallback
            
            # Check if it's actually a CSV file (Anaplan exports CSV with .xlsx extension)
            # Try UTF-8 decoding with BOM handling
            try:
                # Handle UTF-8 BOM
                if file_header.startswith(b'\xef\xbb\xbf'):
                    header_text = file_header[3:].decode('utf-8', errors='ignore')
                else:
                    header_text = file_header.decode('utf-8', errors='ignore')
                
                # Check for CSV characteristics: contains commas/tabs/semicolons, has newlines
                delimiters = [',', '\t', ';']
                has_delimiter = any(d in header_text for d in delimiters)
                has_newlines = '\n' in header_text or '\r' in header_text
                
                if has_delimiter and has_newlines:
                    # Count delimiter occurrences per line for consistency check
                    lines = header_text.replace('\r\n', '\n').replace('\r', '\n').strip().split('\n')[:5]
                    
                    for delimiter in delimiters:
                        counts = [line.count(delimiter) for line in lines if line.strip()]
                        if counts and min(counts) > 0 and max(counts) - min(counts) <= 2:
                            # Consistent delimiter count - likely CSV
                            logger.info(f"Detected CSV content (delimiter: '{delimiter}') in file with Excel extension: {file_path}")
                            logger.info(f"Processing as CSV instead of Excel")
                            return process_csv_file(file_path)
                            
            except Exception as csv_detect_e:
                logger.warning(f"CSV detection failed: {csv_detect_e}")
                
        except Exception as debug_e:
            logger.error(f"Could not check file content: {debug_e}")
        
        return None
    
    # Excel-parsing error patterns that suggest the file isn't a real Excel file
    EXCEL_FORMAT_ERROR_PATTERNS = [
        'not a zip file',
        'invalid zip signature', 
        'does not support file format',
        'file format cannot be determined',
        'unsupported format',
        'bad magic number',
        'not an excel file',
        'zipfile.badzipfile',
        'xlrderror',
        'package not found'
    ]
    
    for engine in engines_to_try:
        try:
            content = try_with_engine(engine)
            if engine:
                logger.info(f"Successfully processed Excel file {file_path} using {engine} engine")
            return content
            
        except Exception as e:
            error_msg = str(e).lower()
            
            # If it's an engine detection error, try the next engine
            if ('engine' in error_msg and 'format cannot be determined' in error_msg) or \
               ('xlrd' in error_msg) or ('openpyxl' in error_msg) or \
               ('no module' in error_msg):
                if engine != engines_to_try[-1]:  # Not the last engine to try
                    logger.warning(f"Engine {engine or 'default'} failed for {file_path}: {e}. Trying next engine...")
                    continue
            
            # Check if error suggests file isn't actually Excel format
            is_format_error = any(pattern in error_msg for pattern in EXCEL_FORMAT_ERROR_PATTERNS)
            
            if is_format_error:
                logger.warning(f"Excel format error detected: {e}")
                # Try CSV fallback for format errors
                csv_result = _try_csv_fallback(file_path)
                if csv_result is not None:
                    return csv_result
            
            logger.error(f"Error processing Excel file {file_path} with engine {engine or 'default'}: {e}")
            return [{'text': f"Error processing Excel file: {str(e)}", 'error': True}]
    
    # This shouldn't be reached, but just in case
    return [{'text': "Error: Unable to process Excel file with any available engine", 'error': True}]

def process_csv_file(file_path: str) -> List[Dict[str, Any]]:
    """Process CSV files and extract structured data with enhanced time-series analysis"""
    try:
        df = pd.read_csv(file_path)
        
        # Create enhanced table text with time-series analysis
        filename = os.path.basename(file_path).replace('.csv', '')
        enhanced_text = _create_enhanced_table_text(df, filename)
        
        # Convert to basic text representation for backward compatibility
        csv_text = df.to_string(index=False)
        
        # Convert pandas data types to Python native types for JSON serialization
        structured_data = []
        if len(df) > 0:
            for record in df.to_dict('records'):
                clean_record = {}
                for key, value in record.items():
                    # Convert numpy types to Python native types
                    if hasattr(value, 'item'):
                        clean_record[key] = value.item()
                    elif pd.isna(value):
                        clean_record[key] = None
                    else:
                        clean_record[key] = value
                structured_data.append(clean_record)
        
        # Perform time-series analysis for CSV data
        time_series_analysis = None
        try:
            from time_series_analyzer import analyze_financial_time_series
            time_series_analysis = analyze_financial_time_series(
                df, 
                metadata={'sheet_name': filename, 'source': 'csv_file', 'file_path': file_path}
            )
            logger.info(f"Time-series analysis completed for CSV: {filename}")
        except ImportError:
            logger.info("Time-series analyzer not available for CSV processing")
        except Exception as e:
            logger.warning(f"Time-series analysis failed for CSV {filename}: {e}")
        
        content = [{
            'text': csv_text,
            'enhanced_text': enhanced_text,  # New: Enhanced text with time-series analysis
            'structured_data': structured_data,
            'columns': list(df.columns),
            'row_count': int(len(df)),  # Convert to Python int
            'file_type': 'csv',
            'time_series_analysis': time_series_analysis,  # New: Include time-series analysis
            'data_quality_score': time_series_analysis.get('summary', {}).get('data_quality_score', 0.8) if time_series_analysis else 0.8,
            'detected_scenarios': time_series_analysis.get('data_normalization', {}).get('scenarios', []) if time_series_analysis else [],
            'detected_metrics': time_series_analysis.get('data_normalization', {}).get('metrics', []) if time_series_analysis else [],
            'fiscal_calendar': time_series_analysis.get('fiscal_calendar', {}) if time_series_analysis else {},
            'has_time_series_data': time_series_analysis is not None,
            'file_path': file_path
        }]
        
        return content
        
    except Exception as e:
        logger.error(f"Error processing CSV file {file_path}: {e}")
        return [{'text': f"Error processing CSV file: {str(e)}", 'error': True}]

def process_text_file(file_path: str, file_format: str) -> List[Dict[str, Any]]:
    """Process text-based files (txt, md, json, xml)"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            raw_content = f.read()
        
        if file_format == 'json':
            # Parse and format JSON for better readability
            try:
                json_data = json.loads(raw_content)
                formatted_content = json.dumps(json_data, indent=2)
                
                return [{
                    'text': formatted_content,
                    'structured_data': json_data if isinstance(json_data, (dict, list)) else {},
                    'file_type': 'json'
                }]
            except json.JSONDecodeError:
                # Fallback to raw text if JSON is invalid
                pass
        
        elif file_format == 'xml':
            # Parse and format XML
            try:
                root = ET.fromstring(raw_content)
                # Convert XML to readable text format
                xml_text = ET.tostring(root, encoding='unicode', method='text')
                
                return [{
                    'text': xml_text,
                    'raw_xml': raw_content,
                    'file_type': 'xml'
                }]
            except ET.ParseError:
                # Fallback to raw text if XML is invalid
                pass
        
        # For txt, md, or invalid json/xml, return as plain text
        return [{
            'text': raw_content,
            'file_type': file_format
        }]
        
    except Exception as e:
        logger.error(f"Error processing text file {file_path}: {e}")
        return [{'text': f"Error processing text file: {str(e)}", 'error': True}]

def process_html_file(file_path: str) -> List[Dict[str, Any]]:
    """Process HTML files and extract readable text"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            html_content = f.read()
        
        # Extract text using custom HTML parser
        parser = HTMLTextExtractor()
        parser.feed(html_content)
        extracted_text = parser.get_text()
        
        return [{
            'text': extracted_text,
            'raw_html': html_content,
            'file_type': 'html'
        }]
        
    except Exception as e:
        logger.error(f"Error processing HTML file {file_path}: {e}")
        return [{'text': f"Error processing HTML file: {str(e)}", 'error': True}]

def process_image_file(file_path: str) -> List[Dict[str, Any]]:
    """Process image files using OCR to extract text"""
    try:
        # Open and process image with OCR
        image = Image.open(file_path)
        
        # Convert to RGB if necessary
        if image.mode not in ('RGB', 'L'):
            image = image.convert('RGB')
        
        # Extract text using OCR
        extracted_text = pytesseract.image_to_string(image)
        
        # Get image metadata
        width, height = image.size
        format_info = image.format
        
        return [{
            'text': extracted_text.strip(),
            'image_info': {
                'width': width,
                'height': height,
                'format': format_info,
                'mode': image.mode
            },
            'file_type': 'image'
        }]
        
    except Exception as e:
        logger.error(f"Error processing image file {file_path}: {e}")
        return [{'text': f"Error processing image file: {str(e)}", 'error': True}]

def process_single_document(file_content: bytes, filename: str) -> Dict[str, Any]:
    """Process a single document of any supported format"""
    file_format = get_file_format(filename)
    
    if file_format == 'unknown':
        return {
            'success': False,
            'error': f"Unsupported file format: {filename}",
            'filename': filename
        }
    
    # Save to temporary file
    with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{filename.split(".")[-1]}') as tmp_file:
        tmp_file.write(file_content)
        temp_path = tmp_file.name
    
    try:
        # Process based on file format
        if file_format == 'pdf':
            content = process_pdf(temp_path)
            total_text = '\n\n'.join([page.get('text', '') for page in content])
            
        elif file_format in ['xlsx', 'xls']:
            content = process_excel_file(temp_path)
            total_text = '\n\n'.join([sheet.get('text', '') for sheet in content])
            
        elif file_format == 'csv':
            content = process_csv_file(temp_path)
            total_text = content[0].get('text', '') if content else ''
            
        elif file_format in ['txt', 'md', 'json', 'xml']:
            content = process_text_file(temp_path, file_format)
            total_text = content[0].get('text', '') if content else ''
            
        elif file_format == 'docx':
            from docx import Document as DocxDocument
            doc = DocxDocument(temp_path)

            # --- Non-table paragraphs ---
            para_lines = [p.text for p in doc.paragraphs if p.text.strip()]

            # --- Table sections: repeat column header at top of every N-row block ---
            # Each section becomes a natural \n\n-separated segment so chunk_text()
            # breaks cleanly between them rather than cutting mid-row.
            TABLE_ROWS_PER_SECTION = 25  # ~25 student rows per chunk ≈ 1 500 chars
            table_sections = []
            for table in doc.tables:
                if not table.rows:
                    continue
                header = ' | '.join(
                    cell.text.strip() for cell in table.rows[0].cells
                )
                data_rows = []
                for row in table.rows[1:]:
                    row_text = ' | '.join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        data_rows.append(row_text)

                # Batch data rows into sections, each prefixed with the header
                for i in range(0, max(len(data_rows), 1), TABLE_ROWS_PER_SECTION):
                    batch = data_rows[i:i + TABLE_ROWS_PER_SECTION]
                    if batch:
                        section_text = header + '\n' + '\n'.join(batch)
                        table_sections.append(section_text)

            # Combine: prose first, then table sections — separated by double newlines
            # so chunk_text() treats each as a natural breakpoint.
            all_sections = []
            if para_lines:
                all_sections.append('\n'.join(para_lines))
            all_sections.extend(table_sections)

            text = '\n\n'.join(all_sections)
            if not text.strip():
                logger.warning(f"No text extracted from Word document: {temp_path}")
            content = [{'text': text, 'page': 1, 'format': 'docx'}]
            total_text = text

        elif file_format == 'html':
            content = process_html_file(temp_path)
            total_text = content[0].get('text', '') if content else ''
            
        elif file_format == 'image':
            content = process_image_file(temp_path)
            total_text = content[0].get('text', '') if content else ''
            
        else:
            return {
                'success': False,
                'error': f"Processing not implemented for format: {file_format}",
                'filename': filename
            }
        
        # Clean up temporary file
        if os.path.exists(temp_path):
            os.unlink(temp_path)
        
        if not content or (isinstance(content, list) and all(item.get('error') for item in content)):
            return {
                'success': False,
                'error': f"Failed to extract content from {filename}",
                'filename': filename
            }
        
        # Extract company information from the document
        company_info = None
        if total_text and len(total_text.strip()) > 50:
            try:
                company_info = extract_company_info(total_text, filename)
                logger.info(f"Company extraction for {filename}: {company_info.get('primary_company', 'None')}")
            except Exception as e:
                logger.error(f"Company extraction failed for {filename}: {e}")

        result = {
            'success': True,
            'filename': filename,
            'file_format': file_format,
            'content': content,
            'total_text': total_text,
            'pages_processed': len(content) if isinstance(content, list) else 1
        }
        
        # Add company information if extracted
        if company_info and company_info.get('primary_company'):
            result['company_info'] = company_info
            
        return result
        
    except Exception as e:
        # Clean up temporary file on error
        if os.path.exists(temp_path):
            os.unlink(temp_path)
        
        logger.error(f"Error processing document {filename}: {e}")
        return {
            'success': False,
            'error': str(e),
            'filename': filename
        }

def process_multiple_documents(files_data: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Process multiple documents and prepare them for analysis"""
    results = []
    successful_documents = []
    failed_documents = []
    
    for file_data in files_data:
        file_content = file_data['content']
        filename = file_data['filename']
        
        logger.info(f"Processing document: {filename}")
        
        # Ensure file_content is bytes
        if isinstance(file_content, str):
            file_content = file_content.encode('utf-8')
        
        result = process_single_document(file_content, filename)
        results.append(result)
        
        if result['success']:
            # Prepare document for vector store
            document_info = {
                'id': f"doc_{len(successful_documents)}_{filename}",
                'name': filename,
                'text': result['total_text'],
                'format': result['file_format'],
                'pages': result['pages_processed'],
                'content': result['content']
            }
            successful_documents.append(document_info)
        else:
            failed_documents.append({
                'filename': filename,
                'error': result['error']
            })
    
    return {
        'success': len(successful_documents) > 0,
        'processed_count': len(successful_documents),
        'failed_count': len(failed_documents),
        'documents': successful_documents,
        'failed_documents': failed_documents,
        'detailed_results': results
    }

class MultiFormatProcessor:
    """Class-based interface for multi-format document processing"""
    
    def __init__(self):
        self.supported_formats = get_supported_formats()
    
    def process_document(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Process a single document from bytes content"""
        # Save to temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix=f"_{filename}") as temp_file:
            temp_file.write(file_content)
            temp_path = temp_file.name
        
        try:
            # Process using existing function  
            result = process_single_document(file_content, filename)
            return result
        finally:
            # Clean up temporary file
            if os.path.exists(temp_path):
                os.unlink(temp_path)
    
    def process_multiple(self, files_data: List[Dict[str, Any]]) -> Dict[str, Any]:
        """Process multiple documents from file data"""
        return process_multiple_documents(files_data)
    
    def validate_format(self, filename: str) -> bool:
        """Validate if file format is supported"""
        return validate_file_format(filename)
    
    def get_formats(self) -> Dict[str, str]:
        """Get supported formats"""
        return self.supported_formats

def get_supported_formats() -> Dict[str, str]:
    """Get dictionary of supported file formats and their descriptions"""
    return {
        'pdf': 'PDF Document',
        'xlsx': 'Excel Spreadsheet (New)',
        'xls': 'Excel Spreadsheet (Legacy)',
        'csv': 'CSV File',
        'txt': 'Text File',
        'md': 'Markdown File',
        'json': 'JSON File',
        'xml': 'XML File',
        'html': 'HTML File',
        'htm': 'HTML File',
        'jpg': 'JPEG Image',
        'jpeg': 'JPEG Image',
        'png': 'PNG Image',
        'bmp': 'BMP Image',
        'gif': 'GIF Image'
    }



if __name__ == "__main__":
    # Test the processor
    print("Multi-Format Document Processor")
    print("Supported formats:", list(get_supported_formats().keys()))